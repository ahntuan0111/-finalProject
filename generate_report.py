import os
import sys
import subprocess

def install_and_import():
    try:
        import docx
    except ImportError:
        print("Installing python-docx...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx
    return docx

docx = install_and_import()

from docx import Document
from docx.shared import Pt
import re

def add_formatted_run(p, text):
    """Parses simple **bold** and *italic* markdown and adds to paragraph."""
    # Split by bold first
    parts_bold = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts_bold:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
        else:
            # Split by italic
            parts_italic = re.split(r'(\*.*?\*)', part)
            for ipart in parts_italic:
                if ipart.startswith('*') and ipart.endswith('*'):
                    run = p.add_run(ipart[1:-1])
                    run.italic = True
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)
                else:
                    run = p.add_run(ipart)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(12)

def main():
    md_file_path = r"c:\Users\Paul\.gemini\antigravity\brain\6142a0f9-cf41-4389-9703-53d9d213741c\report_outline.md.resolved"
    output_dir = r"e:\Final_Project\docx"
    output_path = os.path.join(output_dir, "Report_Outline.docx")
    
    os.makedirs(output_dir, exist_ok=True)
    
    doc = Document()
    
    # Set default style to Times New Roman 12
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    with open(md_file_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
            
        p = doc.add_paragraph()
        p.style = doc.styles['Normal']
        
        # Determine heading level
        if line.startswith("# "):
            text = line[2:]
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
        elif line.startswith("## "):
            text = line[3:]
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
        elif line.startswith("### "):
            text = line[4:]
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
        elif line.startswith("#### "):
            text = line[5:]
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
        else:
            add_formatted_run(p, line)

    doc.save(output_path)
    print(f"Generated {output_path} successfully.")

if __name__ == '__main__':
    main()
