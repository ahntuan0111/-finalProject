import os
import sys

try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def main():
    doc_path = r"e:\Final_Project\docx\MaSci_Final_Report_V11.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    # 1. Delete the orphan Gantt chart table which is currently the first table
    # Let's verify it's the orphan table
    if len(doc.tables) > 0:
        first_table = doc.tables[0]
        try:
            val = first_table.cell(0, 0).text.lower()
            if "task" in val or "phase" in val or "sprint" in val or "no" in val:
                t_elem = first_table._element
                t_elem.getparent().remove(t_elem)
                print("Removed orphan table from Chapter 3.")
        except Exception as e:
            pass

    # 2. Add the table and rename placeholder in 1.7
    chap1_7_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if "1.7 IMPLEMENTATION PLAN" in p.text.strip().upper():
            chap1_7_idx = i
            break
            
    if chap1_7_idx != -1:
        # Find the placeholder
        placeholder_idx = -1
        for i in range(chap1_7_idx, chap1_7_idx + 5):
            if i < len(doc.paragraphs) and "INSERT GANTT CHART" in doc.paragraphs[i].text.upper():
                placeholder_idx = i
                break
                
        if placeholder_idx != -1:
            p_placeholder = doc.paragraphs[placeholder_idx]
            # Replace placeholder text
            p_placeholder.text = "As illustrated in Table 1.1, the iterative phases of the Agile model are allocated explicitly across a structured 16-week development timeline."
            p_placeholder.paragraph_format.first_line_indent = Pt(0)
            
            # Insert title paragraph
            title_p = p_placeholder.insert_paragraph_before()
            title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = title_p.add_run("Table 1.1: Project Implementation Gantt Chart")
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.italic = True
            
            # Create Table
            table = doc.add_table(rows=6, cols=3)
            table.style = 'Table Grid'
            
            headers = ["Phase", "Key Tasks", "Duration (Weeks)"]
            for col_idx, header in enumerate(headers):
                cell = table.cell(0, col_idx)
                cell.text = header
                for paragraph in cell.paragraphs:
                    for run_h in paragraph.runs:
                        run_h.font.name = 'Times New Roman'
                        run_h.font.size = Pt(12)
                        run_h.bold = True
            
            data = [
                ("Phase 1: Analysis & Design", "Requirement scoping, UI/UX mockups, ERD design", "Week 1 - Week 3"),
                ("Phase 2: Core Backend", "Spring Boot Setup, Security JWT, CRUD APIs", "Week 4 - Week 7"),
                ("Phase 3: Frontend Development", "Jetpack Compose Mobile App, React Web Admin", "Week 8 - Week 11"),
                ("Phase 4: AI Integration", "Google Gemini Prompt Engineering, Firebase Storage", "Week 12 - Week 14"),
                ("Phase 5: UAT & Deployment", "User Acceptance Testing, Bug fixing, Final Launch", "Week 15 - Week 16")
            ]
            
            for row_idx, row_data in enumerate(data, start=1):
                for col_idx, cell_data in enumerate(row_data):
                    cell = table.cell(row_idx, col_idx)
                    cell.text = cell_data
                    for paragraph in cell.paragraphs:
                        for run_c in paragraph.runs:
                            run_c.font.name = 'Times New Roman'
                            run_c.font.size = Pt(12)
                            
            # Move the table to right below the title paragraph
            title_p._element.addnext(table._element)

            output_path = r"e:\Final_Project\docx\MaSci_Final_Report_V12.docx"
            doc.save(output_path)
            print("Success! Gantt chart moved to Chapter 1 and saved to " + output_path)
        else:
            print("Could not find the GANTT CHART placeholder in 1.7.")
    else:
        print("Error: Could not find 1.7 Implementation Plan.")
        
if __name__ == '__main__':
    main()
