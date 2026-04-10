import os
import sys

try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def apply_format(p, text, is_bold=False):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = is_bold
        return run
    return None

def insert_image(p_ref, img_path, caption):
    if os.path.exists(img_path):
        new_p = p_ref.insert_paragraph_before()
        new_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        new_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        new_p.paragraph_format.space_after = Pt(0)
        run = new_p.add_run()
        run.add_picture(img_path, width=Inches(6.0))
        
        cap_p = p_ref.insert_paragraph_before()
        cap_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        cap_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap_p.paragraph_format.space_after = Pt(24)
        c_run = cap_p.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(12)
        c_run.italic = True
    else:
        new_p = p_ref.insert_paragraph_before()
        apply_format(new_p, f"[IMAGE NOT FOUND: {img_path}]", is_bold=True)
        new_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

def main():
    doc_path = r"e:\Final_Project\docx\MaSci_Final_Report_V15.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    start_idx = -1
    end_idx = -1
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip().upper()
        if "3.7.1 AI GRADING SEQUENCE DIAGRAM" in text:
            # We want to replace from 3.7.1 down to UI/UX
            start_idx = i
        elif "3.8 UI/UX DESIGN SPECIFICATIONS" in text:
            end_idx = i
            break
            
    if start_idx != -1 and end_idx != -1:
        # Delete old bad text paragraphs
        paragraphs_to_delete = doc.paragraphs[start_idx:end_idx]
        p_ref = doc.paragraphs[end_idx]
        
        for p in paragraphs_to_delete:
            delete_paragraph(p)

        def insert_p(text, bold=False):
            new_p = p_ref.insert_paragraph_before()
            apply_format(new_p, text, is_bold=bold)
            return new_p

        insert_p("3.7.1 AI Grading Sequence Diagram", bold=True)
        insert_p("[1. Business Logic & Purpose]\nThe Sequence Diagram (Figure 3.7) illustrates the most computationally complex processing flow of the system: Assignment Submission and Automated AI Grading. The primary objective of this flow is to replace manual interventions from human educators, providing immediate formative feedback mechanisms for the student's handwritten mathematical exercises.")
        insert_p("[2. Step-by-Step Technical Execution]\nThe execution lifecycle is strictly decoupled to optimize data streams. Initially, the Mobile Client dispatches a synchronous HTTP GET request toward the ExerciseController to solicit upload permissions. The FirebaseStorageService communicates with the Firebase Admin SDK, subsequently returning a time-bound Signed URL. Subsequently, the Mobile Client utilizes this exact URL to execute a PUT command, uploading the image file directly to the Cloud bucket, entirely bypassing the Spring Boot server. After the image is successfully persisted, the Client transmits the resulting URI back to the ExerciseController. Here, the GeminiValidationService is invoked to encapsulate the URI and the mathematical prompt into a Multimodal JSON payload, dispatching an asynchronous HTTP POST request to the Google Gemini API. Finally, upon receiving the structured response, the ExerciseService executes a database transaction to persist the computed score and detailed heuristic reasoning into the exercise_submissions table in MySQL.")
        insert_p("[3. Exception Handling (Fragments)]\nWithin the Alternative Fragment (alt block), the Fault Tolerance mechanism is strictly defined. Should the Google API encounter an outage (Timeout or yielding a 500 Internal Server Error), the GeminiValidationService proactively intercepts this exception. Rather than permitting an abrupt system crash, the service throws a handled exception, explicitly dictating the Controller to return an HTTP 503 (Service Unavailable) to the Client. The Mobile App captures this code and renders a safe dialog, structurally preventing Application Not Responding (ANR) lockups.")
        insert_p("[4. Architectural Justification]\nThe architectural decision to utilize Firebase Signed URLs instead of traditional Multipart File uploads constitutes a core architectural calculation. Pushing massive raw image binaries from thousands of concurrent devices directly through the Spring Boot server would inherently exhaust the Tomcat memory heap, fabricating a systemic bottleneck. Delegating static media storage natively to Firebase ensures the Backend remains entirely Stateless—exclusively parsing lightweight textual attributes (JSON/URIs)—thereby guaranteeing API latency stays mathematically optimal.")
        insert_image(p_ref, r"e:\Final_Project\docx\SequenceDIagram.png", "Figure 3.7: AI Grading Sequence Diagram")

        insert_p("3.7.2 AI Tutor Context Sequence Diagram", bold=True)
        insert_p("[1. Business Logic & Purpose]\nThe AI Tutor Context Sequence Diagram (Figure 3.8) depicts the foundational execution cycle driving the Socratic Conversational AI. Rather than offering immediate answers, this conversational flow allows the AI to interrogate the student continuously, dynamically adapting based upon the student's mathematical comprehension.")
        insert_p("[2. Step-by-Step Technical Execution]\nThe lifecycle originates within an Optional Fragment (opt block), where the AiTutorService determines if the user is establishing a new conversation or resuming an existing thread by evaluating the ai_chat_sessions table. Inside the dominant Loop Fragment, the fundamental Context Injection technique occurs. The Backend explicitly queries the ai_chat_messages table from MySQL to retrieve historical conversations. It rigorously segregates these payloads strictly by the role enum ('USER' or 'ASSISTANT'), injecting them chronologically into a consolidated JSON array. This dense contextual array is transmitted to the Gemini API, ensuring the Large Language Model comprehensively understands the exact state of the ongoing mathematical discourse before generating its next hint.")
        insert_p("[3. Exception Handling (Fragments)]\nSimilar to the Grading Flow, the Alternative (alt) Fragment securely manages external timeout anomalies. If the external LLM boundary breaches standard execution vectors (e.g., exceeding connection timeouts), the AiTutorService actively intercepts the disruption. It halts the parsing loop and issues an HTTP 503 fallback exception, enabling the Mobile UI to cleanly reset the chat interface instead of perpetually freezing the user's application thread.")
        insert_p("[4. Architectural Justification]\nThe explicit decision to persist conversational trajectories inside relational MySQL tables (ai_chat_sessions and ai_chat_messages) instead of leveraging volatile in-memory databases like Redis, or localized mobile data stores like SQLite, is a calculated necessity. This centralized persistence architecture guarantees absolute cross-platform continuity. Should a student switch physical devices or migrate their learning session toward the Web Administration portal, the entire historical Socratic flow is preserved and instantly traversable.")
        insert_image(p_ref, r"e:\Final_Project\docx\SequenceDiagram (2).png", "Figure 3.8: AI Tutor Context Sequence")

        output_path = r"e:\Final_Project\docx\MaSci_Final_Report_V16.docx"
        doc.save(output_path)
        print("Success! Sequence diagrams rewritten to academic perfection in V16: " + output_path)
    else:
        print("Error: Could not find Sequence Diagrams boundaries.")

if __name__ == '__main__':
    main()
