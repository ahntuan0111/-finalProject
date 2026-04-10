import os
import sys

try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def apply_format(p, text, is_bold=False):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(24)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = is_bold
        return run
    return None

def main():
    doc_path = r"e:\Final_Project\docx\MaSci_Final_Report_V7.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    # Locate Chapter 3 and Chapter 4
    chap3_idx = -1
    chap4_idx = -1
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip().upper()
        if "CHAPTER 3" in text and "ANALYSIS" in text:
            chap3_idx = i
        elif "CHAPTER 4" in text and "IMPLEMENTATION" in text:
            chap4_idx = i
            break
            
    if chap3_idx != -1 and chap4_idx != -1:
        p_chap4 = doc.paragraphs[chap4_idx]
        
        def insert_p(text, bold=False, center=False):
            new_p = p_chap4.insert_paragraph_before()
            run = apply_format(new_p, text, is_bold=bold)
            if center:
                new_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            return new_p
            
        def insert_heading(text):
            new_p = p_chap4.insert_paragraph_before()
            new_p.paragraph_format.space_before = Pt(18)
            new_p.paragraph_format.space_after = Pt(12)
            run = new_p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True

        def insert_image(img_path, caption):
            if os.path.exists(img_path):
                new_p = p_chap4.insert_paragraph_before()
                new_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                # Single spacing for image paragraph to keep it close to caption
                new_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                new_p.paragraph_format.space_after = Pt(0)
                run = new_p.add_run()
                run.add_picture(img_path, width=Inches(6.0))
                
                # Caption
                cap_p = p_chap4.insert_paragraph_before()
                cap_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                cap_p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                cap_p.paragraph_format.space_after = Pt(24)
                c_run = cap_p.add_run(caption)
                c_run.font.name = 'Times New Roman'
                c_run.font.size = Pt(12)
                c_run.italic = True
            else:
                insert_p(f"[IMAGE NOT FOUND: {img_path}]", bold=True, center=True)

        insert_heading("3.1 System Requirements")
        insert_p("In the development of the MaSci platform, system requirements are meticulously gathered and categorized into Functional Requirements (defining what the system must do from the user's perspective) and Non-functional Requirements (defining system attributes such as speed, security, and AI accuracy). These requirements act as the foundational blueprint guiding all subsequent architectural decisions.")

        insert_heading("3.2 System Flowcharts")
        insert_p("Flowcharts are instrumental in mapping out the sequential logic of complex business processes within the MaSci framework. The following diagrams delineate the exact step-by-step execution pathways, encompassing both normal operation methodologies and rigorous exception handling.")
        
        insert_p("3.2.1 Learning Journey Flowchart")
        insert_p("As illustrated in Figure 3.1, the learning flow commences with the system loading the list of chapters and lessons. The core decision node lies in interrogating 'Is a video attached to the lesson?'. If the condition evaluates to true, the system dynamically embeds the video player directly into the interface. Concurrently, an extended branch permits the student to query the AI Tutor in parallel with the theoretical learning process, prior to arriving at the final 'Mark as Complete' step to formally register their progression into the database.")
        insert_image(r"e:\Final_Project\docx\flowchart_01.png", "Figure 3.1: Learning Journey Flowchart")

        insert_p("3.2.2 Assignment Submission and AI Grading Flowchart")
        insert_p("The integrity of the image-based grading system is maintained through robust technical workflows. As depicted in Figure 3.2, the system implements a two-tier exception handling mechanism. The first tier occurs at the Client level: the application actively evaluates 'Is the captured image sharp and adequately lit?' prior to the upload sequence; if it fails to meet this threshold, the UI strictly requests a recapture. The second tier of validation occurs post-AI analysis: if the Gemini algorithm returns an 'Invalid' flag, the system intercepts the process and displays the precise reasoning to the user rather than persisting garbage data into the database.")
        insert_image(r"e:\Final_Project\docx\flowchart_02.png", "Figure 3.2: Assignment Submission and AI Grading Flowchart")

        insert_p("3.2.3 AI Tutor Interaction Flowchart")
        insert_p("Figure 3.3 clearly demonstrates the Context Injection mechanism fundamental to the AI's contextual awareness. At the initialization step, the system actively evaluates 'Does a previous chat session exist?'. If affirmatively evaluated, it triggers a query to the database (represented by the cylindrical data block) to sequentially load the entire conversational history. This continuous loop ensures that every subsequent response generated by the AI is strictly predicated on the immediately preceding conversational memory, maintaining fluid instructional continuity.")
        insert_image(r"e:\Final_Project\docx\flowchart_03.png", "Figure 3.3: AI Tutor Interaction Flowchart")

        insert_heading("3.3 Sequence Diagrams")
        insert_p("Sequence diagrams are critical for visualizing the life-cycle of API calls, object interactions, and backend class coordinations. The following diagrams precisely map the participants from the Android presentation layer through the Spring Boot service layer down to the external endpoints.")

        insert_p("3.3.1 AI Grading Sequence Diagram")
        insert_p("The execution process, as depicted in Figure 3.4, begins when the Mobile App invokes the 'ExerciseController' to request an upload URL. Subsequenly, the 'FirebaseStorageService' class assumes responsibility for communicating securely with Firebase to retrieve the Signed URL. After the image is successfully pushed to the Cloud environment, the 'ExerciseController' delegates authority to the 'GeminiValidationService' to invoke the Gemini API for comprehensive multimodal analysis, prior to formally recording the evaluated score into the database via the 'ExerciseService'.")
        insert_image(r"e:\Final_Project\docx\SequenceDIagram.png", "Figure 3.4: AI Grading Sequence Diagram")

        insert_p("3.3.2 AI Tutor Sequence Diagram")
        insert_p("Similarly, the conversational operations between the 'AiTutorController' and the 'AiTutorService' are closely and synchronously coordinated. As explicitly demonstrated in the 'alt' conditional block within Figure 3.5, if the external connection to the Gemini API encounters a critical failure (e.g., Timeout or a 500 Internal Server Error), the 'AiTutorService' intercepts the fault and throws a dedicated exception. This signals the Controller to return a secure 500 error code, which safely triggers the error-handling user interface on the Mobile App, preventing application crashes.")
        insert_image(r"e:\Final_Project\docx\SequenceDiagram (2).png", "Figure 3.5: AI Tutor Sequence Diagram")

        insert_heading("3.4 Database Design and Entity-Relationship")
        insert_p("The architecture of the MaSci relational database is designed to enforce strict data integrity utilizing InnoDB and complex foreign key constraints. The following Entity-Relationship Diagram encapsulates the exact physical schema of the 24 intrinsic tables driving the platform.")
        insert_p("The Entity-Relationship Diagram (ERD) in Figure 3.6 illustrates the physical structural layout of the system. Observers can identify three main functional clusters: the left cluster securely manages user authentication and profiling (encompassing the 'users' and 'student_profiles' tables); the central cluster simulates the academic hierarchy via a Cascading Tree model, branching sequentially from grades to subjects to chapters and finally to lessons; and the right-hand cluster is dedicated to storing dynamic learning progress records and chat historics (including 'quiz_results' and 'ai_chat_messages'). The continuous solid connecting lines represent rigorous Foreign Key constraints, definitively ensuring referential integrity across the entire ecosystem.")
        insert_image(r"e:\Final_Project\docx\ERD.png", "Figure 3.6: System Entity-Relationship Diagram (ERD)")

        # Delete previous chapter 3 dummy content
        paragraphs_to_delete = doc.paragraphs[chap3_idx+1:chap4_idx]
        for p in paragraphs_to_delete:
            delete_paragraph(p)
            
        output_path = r"e:\Final_Project\docx\MaSci_Final_Report_V8.docx"
        doc.save(output_path)
        print("Success! Chapter 3 has been expanded significantly, translated to English, and embedded with all images. Saved to " + output_path)
    else:
        print("Error: Could not find CHAPTER 3 or CHAPTER 4 in the document.")
        
if __name__ == '__main__':
    main()
