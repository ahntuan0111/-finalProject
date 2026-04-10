import os
import sys

try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING

def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None

def apply_format(p, text, is_bold=False):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(24)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = is_bold
    return run

def main():
    doc_path = r"e:\Final_Project\docx\MaSci_Final_Report_V5.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    # Locate Chapter 1 and Chapter 2
    chap1_idx = -1
    chap2_idx = -1
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip().upper()
        if "CHAPTER 1" in text and "INTRODUCTION" in text:
            chap1_idx = i
        elif "CHAPTER 2" in text and ("LITERATURE REVIEW" in text or "THEORETICAL" in text):
            chap2_idx = i
            break
            
    if chap1_idx != -1 and chap2_idx != -1:
        p_chap2 = doc.paragraphs[chap2_idx]
        
        def insert_p(text, bold=False):
            new_p = p_chap2.insert_paragraph_before()
            apply_format(new_p, text, is_bold=bold)
            return new_p
            
        def insert_heading(text):
            new_p = p_chap2.insert_paragraph_before()
            new_p.paragraph_format.space_before = Pt(18)
            new_p.paragraph_format.space_after = Pt(12)
            run = new_p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True

        insert_heading("1.1 Background to the Study")
        insert_p("Following the COVID-19 pandemic, the educational landscape in Vietnam, as well as globally, experienced a substantial paradigm shift towards digital platforms. However, this accelerated transition to online learning has inadvertently illuminated and widened the disparity in educational quality between urban centers and rural areas. Schools located in remote or economically disadvantaged regions are confronting a severe shortage of highly qualified and specialized teachers, particularly in natural science subjects such as Mathematics, Physics, and Chemistry. Consequently, secondary school students (grades 6 to 9) in these areas often struggle with self-studying complex concepts without adequate guidance, facing an acute lack of personalized learning resources.")
        insert_p("In the traditional educational setting, teachers are frequently overburdened with administrative tasks and the manual grading of hundreds of student assignments, which leads to burnout and restricts the time available for individualized student mentorship. The traditional 'one-size-fits-all' pedagogical approach is proving increasingly ineffective as it fails to address the unique learning pace and comprehension level of each student.")
        insert_p("Simultaneously, the global technological sphere has witnessed a transformative boom in Artificial Intelligence (AI), particularly with the advent of Generative AI and Large Language Models (LLMs) such as ChatGPT and Google Gemini throughout 2023 and 2024. These advancements have created unprecedented opportunities for embedding AI-assisted learning into the core educational framework. Specifically, the emergence of Multimodal AI (capable of processing text, voice, and images concurrently) holds the immense potential to completely revolutionize traditional learning methodologies. By leveraging these cutting-edge technologies, the MaSci E-Learning platform is proposed as an urgent, innovative, and comprehensive solution to decentralize quality education, providing secondary school students with equitable access to intelligent, on-demand academic support.")

        insert_heading("1.2 Problem Statement")
        insert_p("Despite the proliferation of online educational platforms, the current state of digital education for secondary school students continues to suffer from several critical unresolved pain points:")
        insert_p("Lack of 1-on-1 Interaction: The most significant drawback of self-study in existing online environments is the absence of immediate, personalized support. When students encounter a difficult mathematical problem or a complex scientific concept, traditional platforms merely offer static text or pre-recorded videos. These systems cannot act as a dedicated tutor to interactually analyze the student's specific misunderstanding, ask guiding questions, or suggest the correct step-by-step approach, leaving students frustrated and unmotivated.")
        insert_p("Technological Barriers in Entering Mathematical Formulas: A major hurdle in digital science education is the complexity of inputting mathematical symbols via standard mobile or computer keyboards. Expecting secondary school students to learn and memorize complex syntaxes, such as LaTeX, shifts their cognitive load from focusing on mathematical reasoning to struggling with technological interfaces. This barrier makes students highly reluctant to submit assignments online. Consequently, capturing a photographic image of handwritten work becomes the most natural and optimal solution for students; however, traditional E-Learning systems lack the capability to automatically process and grade such visual inputs.")
        insert_p("Limitations in Grading Methodology: Currently, most E-Learning systems rely heavily on multiple-choice questions for automated grading due to their technical simplicity. However, multiple-choice formats fundamentally fail to evaluate the step-by-step logical methodology and critical thinking process of a student. While handwritten essay assignments are pedagogically superior for assessing deep comprehension, they cannot be efficiently administered remotely because teachers simply do not possess the time required to manually review, decipher, and score hundreds of submitted exercise images.")
        insert_p("Disjointed Content Management: From an administrative perspective, the management of learning resources is highly fragmented. Teachers expend an enormous amount of time and effort attempting to create, format, and update questions that contain complex molecular structures or mathematical formulas. The lack of a cohesive, user-friendly Web Admin dashboard severely hampers the continuous maintenance and scalability of current educational content.")

        insert_heading("1.3 Research Questions")
        insert_p("In order to construct a robust foundation for the system's development and evaluate its effectiveness, this research addresses the following core questions:")
        insert_p("1. How can a comprehensive 3-tier architecture (encompassing an Android mobile application, a React-based Web Admin, and a Spring Boot Backend) operate seamlessly and be optimized for a high-traffic E-learning platform?")
        insert_p("2. What are the specific pedagogical and operational benefits of integrating Google's Gemini Multimodal AI for accurately grading and analyzing handwritten mathematical exercises?")
        insert_p("3. How does the implementation of an automated progress tracking system (gamification streak) and a Socratic AI Tutor effectively improve the self-study motivation and academic perseverance of secondary school students?")

        insert_heading("1.4 Objectives of the Study")
        insert_p("Primary Objective:")
        insert_p("The overarching goal of this project is to architect, develop, and successfully deploy the MaSci E-Learning platform—a fully comprehensive educational system that integrates advanced Artificial Intelligence to solve the inherent challenges of remote science education in Vietnam.")
        insert_p("Specific Objectives:")
        insert_p("- To engineer a secure, highly scalable Backend Server utilizing Spring Boot 3.2.5 and MySQL to handle data processing and user authentication securely via JWT.")
        insert_p("- To develop a high-performance Web Admin portal using React 19, enabling teachers to effortlessly digitize multimedia theory and organize curriculum structures.")
        insert_p("- To design an intuitive, responsive Mobile Application for the Android platform using Kotlin and Jetpack Compose, granting students seamless access to lectures and quizzes.")
        insert_p("- To integrate Gemini AI to function as a Socratic AI Tutor capable of providing context-aware conversational feedback, and to leverage its multimodal capabilities to automatically grade submitted images of handwritten assignments.")

        insert_heading("1.5 Significance of the Study")
        insert_p("The successful implementation of the MaSci platform brings profound and measurable impacts across four primary vectors:")
        insert_p("- For Students: The platform essentially democratizes access to private tutoring. It provides a 24/7 AI tutor, personalizes learning trajectories, and instantly analyzes handwritten methodology to deliver immediate, constructive feedback, thereby dramatically enhancing self-study efficiency.")
        insert_p("- For Teachers and Administrators: It radically reduces the exhaustive administrative workload associated with manually grading essay-based assignments. Furthermore, it provides a standardized, rich-text authoring system that streamlines the digitization of complex scientific content.")
        insert_p("- For Society: On a macro level, the system acts as a catalyst for educational equity. It substantially narrows the gap in access to high-quality educational resources between metropolitan centers and rural or underprivileged geographical regions.")
        insert_p("- For the Technical Community: The project serves as an empirical case study affirming the practical viability and high performance of synthesizing modern architectural frameworks (Spring Boot, Jetpack Compose, React) with state-of-the-art Generative AI APIs in a production-ready application.")

        insert_heading("1.6 Scope and Limitation")
        insert_p("Scope:")
        insert_p("The scope of this study is strategically focused on secondary school students (specifically grades 6 through 9) studying within the Vietnamese educational curriculum, with an emphasis on natural science subjects. On the technological front, the student-facing platform is exclusively built for the Android operating system using Kotlin and declarative UI via Jetpack Compose. The administrative interface is constructed as a React 19 web application. The core logic and AI API integrations are managed exclusively through a Spring Boot backend infrastructure.")
        insert_p("Limitations:")
        insert_p("Despite its robust architecture, the current iteration of the MaSci system possesses certain limitations. First, the mobile application does not currently support the iOS (Apple) platform, restricting access for a segment of smartphone users. Second, while the AI integration is highly advanced, grading accuracy may still be susceptible to a margin of error if the student's uploaded image is exceedingly blurry or if the handwriting is exceptionally illegible. Finally, the system currently lacks integrations with third-party online Payment Gateways, restricting the immediate deployment of premium monetization models.")

        insert_heading("1.7 Implementation Plan")
        insert_p("The MaSci software development lifecycle is strictly executed following the Agile Scrum methodology across a comprehensive 16-week timeline. This iterative approach ensures flexibility, continuous integration, and consistent evaluation against core educational requirements. The development process is meticulously divided into distinct time-boxed Sprints. Initial Sprints are dedicated to rigorous Requirements Gathering and Architectural Design. Subsequent Sprints encompass the Backend Implementation, Database orchestration, and Web Admin development. Final Sprints focus intensely on cutting-edge AI Integration, Mobile UI construction, and exhaustive User Acceptance Testing (UAT).")
        insert_p("[INSERT GANTT CHART IMAGE HERE: A detailed Gantt chart illustrating the 16-week timeline and milestone achievements is visually presented in Chapter 3.]")

        insert_heading("1.8 Definition of Terms")
        insert_p("To ensure absolute clarity and optimize the academic reading experience, the following technical terminologies utilized throughout this report are defined:")
        insert_p("- Multimodal AI: Advanced artificial intelligence models that possess the capability to simultaneously ingest, understand, and synthesize multiple forms of input data modalities—such as text, visual images, and auditory signals—to generate cohesive responses (e.g., Google Gemini).")
        insert_p("- Socratic Method: A pedagogical approach predicated on asking continuous, open-ended, and thought-provoking questions. Within MaSci, the AI Tutor employs this method to skillfully guide students toward discovering the correct answers conceptually, rather than outright providing the final mathematical solution.")
        insert_p("- JWT (JSON Web Token): A compact, URL-safe, open standard utilized for securely transmitting information and claims between two parties. In this architecture, it is fundamentally used to authenticate and authorize user sessions.")
        insert_p("- Jetpack Compose: An innovative, modern declarative user interface toolkit provided by Google, designed to accelerate and simplify native UI development for Android applications.")
        insert_p("- Single Page Application (SPA): An architectural pattern for web applications wherein a single HTML document is loaded, and subsequent interface components are dynamically retrieved and updated by JavaScript. This ensures a highly fluid user experience for the React Admin dashboard.")
        insert_p("- Gamification: The strategic application of game-design elements, such as the 'Streak' progress tracking system implemented in MaSci, to a non-game context in order to maximize user engagement, competitiveness, and sustained learning motivation.")
        insert_p("- MVVM (Model-View-ViewModel): A structural design pattern utilized in the Android application to facilitate a clear separation of graphical user interface development from the underlying business logic and data models.")
        insert_p("- DTO (Data Transfer Object): An object specifically utilized to encapsulate and transport data across various subsystems or network layers (e.g., between the Spring Boot Backend and the front-end clients), ensuring secure and optimized data handling.")

        # Identify previously added chapter 1 items and delete them
        paragraphs_to_delete = doc.paragraphs[chap1_idx+1:chap2_idx]
        for p in paragraphs_to_delete:
            delete_paragraph(p)
            
        output_path = r"e:\Final_Project\docx\MaSci_Final_Report_V6.docx"
        doc.save(output_path)
        # Safe printing without unicode characters
        print("Success! Chapter 1 has been expanded significantly and saved to " + output_path)
    else:
        print("Error: Could not find CHAPTER 1 or CHAPTER 2 in the document.")
        
if __name__ == '__main__':
    main()
