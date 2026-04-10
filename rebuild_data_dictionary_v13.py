import os
import sys

try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING

def apply_format(p, text, is_bold=False, is_italic=False):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = is_bold
        run.italic = is_italic
        return run
    return None

def create_db_table(document, p_ref, table_title, table_desc, headers, rows_data):
    # Insert Title
    title_p = p_ref.insert_paragraph_before()
    title_p.paragraph_format.space_before = Pt(12)
    run_t = title_p.add_run(table_title)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(12)
    run_t.bold = True

    # Insert Description
    desc_p = p_ref.insert_paragraph_before()
    desc_p.paragraph_format.space_after = Pt(6)
    run_d = desc_p.add_run(f"({table_desc})")
    run_d.font.name = 'Times New Roman'
    run_d.font.size = Pt(11)
    run_d.italic = True

    # Create Table
    table = document.add_table(rows=len(rows_data)+1, cols=len(headers))
    table.style = 'Table Grid'
    
    # Fill headers
    for col_idx, header in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header
        for paragraph in cell.paragraphs:
            for run_h in paragraph.runs:
                run_h.font.name = 'Times New Roman'
                run_h.font.size = Pt(11)
                run_h.bold = True
                
    # Fill data
    for row_idx, row_data in enumerate(rows_data, start=1):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_value)
            for paragraph in cell.paragraphs:
                for run_c in paragraph.runs:
                    run_c.font.name = 'Times New Roman'
                    run_c.font.size = Pt(11)

    # Move table before the reference paragraph
    p_ref._element.addprevious(table._element)
    
    # Add spacing after table
    sp_p = p_ref.insert_paragraph_before()
    sp_p.paragraph_format.space_after = Pt(12)

def main():
    doc_path = r"e:\Final_Project\docx\MaSci_Final_Report_V12.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    start_idx = -1
    end_idx = -1
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip().upper()
        if "3.5.1 DATA DICTIONARY" in text or "3.5.2 DATA DICTIONARY" in text:
            start_idx = i
        elif "3.6 SYSTEM FLOWCHARTS" in text:
            end_idx = i
            break
            
    if start_idx != -1 and end_idx != -1:
        # Delete old dictionary paragraphs
        paragraphs_to_delete = doc.paragraphs[start_idx:end_idx]
        p_ref = doc.paragraphs[end_idx]
        
        for p in paragraphs_to_delete:
            p_elem = p._element
            p_elem.getparent().remove(p_elem)
            p._p = p._element = None
            
        def insert_heading(text, level=2):
            new_p = p_ref.insert_paragraph_before()
            new_p.paragraph_format.space_before = Pt(18)
            new_p.paragraph_format.space_after = Pt(12)
            run = new_p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14 if level == 1 else 12)
            run.bold = True

        insert_heading("3.5.2 Data Dictionary", level=3)
        
        intro_p = p_ref.insert_paragraph_before()
        apply_format(intro_p, "To optimize the presentation and highlight the core business flows of the MaSci platform, the following Data Dictionary strictly focuses on analyzing the fundamental entity tables that drive the student's learning journey and the Artificial Intelligence functionalities. The comprehensive structural schema of all 24 system tables is fully attached in Appendix B of this report.")
        
        headers = ["Field Name", "Data Type", "Constraint", "Description"]

        # Group 1
        g1_p = p_ref.insert_paragraph_before()
        apply_format(g1_p, "Group 1: Academic Content Organization", is_bold=True)
        
        create_db_table(doc, p_ref, "Table 3.2: Structure of 'lessons' table", 
            "Description: This table represents the core atomic unit of the academic structure, establishing strict dependencies with parent chapters.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["chapter_id", "BIGINT", "FK (chapters.id)", "Reference to parent chapter"],
            ["title", "VARCHAR", "NOT NULL", "Standardized title of the lesson"],
            ["sequence_order", "INT", "NOT NULL", "Display order hierarchy"]
        ])
        
        create_db_table(doc, p_ref, "Table 3.3: Structure of 'lesson_content' table", 
            "Description: Crucial for rendering diverse pedagogies. The content_type field explicitly dictates whether the Client UI should load a YouTube Video instance or a native Rich Text parser.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["lesson_id", "BIGINT", "FK (lessons.id)", "Target lesson"],
            ["content_type", "ENUM", "NOT NULL", "'VIDEO' or 'RICH_TEXT'"],
            ["content_body", "TEXT", "NULL", "Markdown/KaTeX theoretical body"],
            ["video_url", "VARCHAR", "NULL", "External multimedia source URL"]
        ])

        # Group 2
        g2_p = p_ref.insert_paragraph_before()
        apply_format(g2_p, "Group 2: Exercise & AI Validation Flow", is_bold=True)
        
        create_db_table(doc, p_ref, "Table 3.4: Structure of 'exercises' table", 
            "Description: Stores mathematical problems alongside scaffolding hints designed to prevent cognitive overload during independent study.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["chapter_id", "BIGINT", "FK (chapters.id)", "Target chapter context"],
            ["prompt_content", "TEXT", "NOT NULL", "The actual mathematical prompt"],
            ["hint", "TEXT", "NULL", "Optional scaffolding guidance"]
        ])
        
        create_db_table(doc, p_ref, "Table 3.5: Structure of 'exercise_submissions' table", 
            "Description: This table serves as the epicenter of the AI Grading capabilities. The image_url explicitly persists the Firebase Signed URL, whereas the score and ai_reasoning fields are updated asynchronously subsequent to evaluating constraints from the Google Gemini API.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["user_id", "BIGINT", "FK (users.id)", "Student invoking submission"],
            ["exercise_id", "BIGINT", "FK (exercises.id)", "Targeted exercise ID"],
            ["image_url", "TEXT", "NOT NULL", "Firebase Storage secured URL"],
            ["score", "DOUBLE", "NULL", "Computed AI evaluation score"],
            ["ai_reasoning", "TEXT", "NULL", "Detailed heuristic assessment from Gemini"]
        ])

        # Group 3
        g3_p = p_ref.insert_paragraph_before()
        apply_format(g3_p, "Group 3: Quiz & Progress Tracking", is_bold=True)
        
        create_db_table(doc, p_ref, "Table 3.6: Structure of 'quizzes' table", 
            "Description: Facilitates traditional multiple-choice evaluations linked precisely to specific hierarchical lessons.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["lesson_id", "BIGINT", "FK (lessons.id)", "Correlated lesson context"],
            ["title", "VARCHAR", "NOT NULL", "Assessment title"],
            ["total_questions", "INT", "NOT NULL", "Count of aggregated questions"]
        ])
        
        create_db_table(doc, p_ref, "Table 3.7: Structure of 'quiz_results' table", 
            "Description: Optimizes query execution arrays by denormalizing holistic metrics. Persisting the final percentage score accelerates gamification telemetry calculations.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["user_id", "BIGINT", "FK (users.id)", "Executing learner"],
            ["quiz_id", "BIGINT", "FK (quizzes.id)", "Assessment target"],
            ["total_score", "DOUBLE", "NOT NULL", "Post-calculation percentage"],
            ["timestamp", "DATETIME", "DEFAULT CURRENT", "Chronological stamp"]
        ])
        
        create_db_table(doc, p_ref, "Table 3.8: Structure of 'lesson_completion' table", 
            "Description: The foundational construct powering analytical progression maps and resolving sequential learning prerequisites (protecting future modules until completion is toggled).", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["user_id", "BIGINT", "FK (users.id)", "Evaluating learner"],
            ["lesson_id", "BIGINT", "FK (lessons.id)", "Specific theoretical block"],
            ["is_completed", "BOOLEAN", "DEFAULT FALSE", "Binary completion gate"]
        ])

        # Group 4
        g4_p = p_ref.insert_paragraph_before()
        apply_format(g4_p, "Group 4: Conversational AI Environment", is_bold=True)
        
        create_db_table(doc, p_ref, "Table 3.9: Structure of 'ai_chat_messages' table", 
            "Description: The backbone array driving the continuous AI Tutor Socratic subsystem. The fundamental 'role' attribute plays a critical isolation role, enabling the Backend architecture to accurately serialize the conversational Context Window JSON representation before routing data toward the external LLM.", headers, [
            ["id", "BIGINT", "PK, AUTO_INCREMENT", "Unique identifier"],
            ["session_id", "BIGINT", "FK (ai_chat_sessions.id)", "Master conversation chain"],
            ["role", "ENUM", "NOT NULL", "'USER' or 'ASSISTANT' designation"],
            ["content", "TEXT", "NOT NULL", "Rich conversational string mapping"],
            ["created_at", "DATETIME", "DEFAULT CURRENT", "Timestamp vector mapping"]
        ])

        output_path = r"e:\Final_Project\docx\MaSci_Final_Report_V13.docx"
        doc.save(output_path)
        print("Success! Data dictionary pruned and rebuilt to V13: " + output_path)
    else:
        print("Error: Could not find Data Dictionary boundaries.")

if __name__ == '__main__':
    main()
